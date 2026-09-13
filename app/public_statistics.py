from __future__ import annotations

import json
import logging
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


logger = logging.getLogger(__name__)
JST = ZoneInfo("Asia/Tokyo")
JNTO_PAGE = "https://www.jnto.go.jp/statistics/data/visitors-statistics/"
CONSUMPTION_PAGE = "https://www.mlit.go.jp/kankocho/tokei_hakusyo/gaikokujinshohidoko.html"
LAB_INBOUND_URL = "https://lab.ugatta-llc.com/data/inbound/latest.json"
LAB_ANALYSIS_URL = "https://lab.ugatta-llc.com/data/analysis/dx-necessity.json"
ALLOWED_HOSTS = {"www.jnto.go.jp", "www.mlit.go.jp", "lab.ugatta-llc.com"}
MAX_EXCEL_BYTES = 50 * 1024 * 1024


def _normal(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def _published_from_url(url: str) -> str | None:
    match = re.search(r"/(20\d{2})(\d{2})(\d{2})_", url)
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}" if match else None


def _published_from_page(value: str) -> str | None:
    match = re.search(r"(20\d{2})年\s*(\d{1,2})月\s*(\d{1,2})日", _normal(value))
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}" if match else None


def _get(session: requests.Session, url: str, *, timeout: int = 60) -> requests.Response:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in ALLOWED_HOSTS:
        raise ValueError(f"許可されていない取得先です: {url}")
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return response


def _get_excel(session: requests.Session, url: str) -> bytes:
    response = _get(session, url, timeout=90)
    content_type = response.headers.get("Content-Type", "").lower()
    if not any(label in content_type for label in ("spreadsheet", "excel", "octet-stream")):
        raise RuntimeError(f"ExcelではないContent-Typeです: {content_type or '未記録'}")
    content = response.content
    if len(content) > MAX_EXCEL_BYTES:
        raise RuntimeError("Excelファイルがサイズ上限を超えています。")
    if not content.startswith(b"PK"):
        raise RuntimeError("Excelファイルのシグネチャが不正です。")
    return content


def _jnto_statistics(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    response = _get(session, JNTO_PAGE)
    response.encoding = response.apparent_encoding or response.encoding
    soup = BeautifulSoup(response.text, "html.parser")
    candidates = []
    for anchor in soup.find_all("a", href=True):
        text = _normal(anchor.get_text(" ", strip=True))
        href = urljoin(JNTO_PAGE, anchor["href"])
        if "国籍/月別" in text and "訪日外客数" in text and href.lower().endswith(".xlsx"):
            candidates.append(href)
    if not candidates:
        raise RuntimeError("JNTOの時系列Excelを特定できませんでした。")
    excel_url = candidates[0]
    workbook = load_workbook(BytesIO(_get_excel(session, excel_url)), read_only=True, data_only=True)
    try:
        year_sheets = [sheet for sheet in workbook.worksheets if re.fullmatch(r"20\d{2}", sheet.title)]
        if not year_sheets:
            raise RuntimeError("JNTO Excelの年別シートを特定できませんでした。")
        sheet = max(year_sheets, key=lambda item: int(item.title))
        header = [cell.value for cell in sheet[4]]
        total = [cell.value for cell in sheet[5]]
        records: list[tuple[int, int, float | None]] = []
        for index, label in enumerate(header):
            match = re.fullmatch(r"\s*(\d{1,2})月\s*", _normal(label))
            if not match or index >= len(total) or not isinstance(total[index], (int, float)):
                continue
            yoy = total[index + 1] if index + 1 < len(total) and isinstance(total[index + 1], (int, float)) else None
            records.append((int(match.group(1)), int(total[index]), float(yoy) if yoy is not None else None))
        if not records:
            raise RuntimeError("JNTO Excelから全国値を抽出できませんでした。")
        month, value, yoy = records[-1]
        period = f"{int(sheet.title):04d}-{month:02d}"
    finally:
        workbook.close()
    source = {
        "source_name": "JNTO「訪日外客統計」",
        "source_url": JNTO_PAGE,
        "excel_url": excel_url,
        "reference_period": period,
        "published_date": _published_from_url(excel_url),
        "release_type": "推計値・暫定値を含む時系列表",
    }
    indicator = {"value": value, "unit": "人", "yoy_pct": yoy, "source_key": "jnto_visitors"}
    return source, indicator


def _consumption_statistics(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    response = _get(session, CONSUMPTION_PAGE)
    response.encoding = response.apparent_encoding or response.encoding
    soup = BeautifulSoup(response.text, "html.parser")
    chosen = None
    for anchor in soup.find_all("a", href=True):
        text = _normal(anchor.get_text(" ", strip=True))
        href = urljoin(CONSUMPTION_PAGE, anchor["href"])
        if "集計表" in text and "速報" in text and href.lower().endswith(".xlsx"):
            chosen = (text, href)
            break
    if chosen is None:
        raise RuntimeError("インバウンド消費動向調査の最新Excelを特定できませんでした。")
    label, excel_url = chosen
    page_text = _normal(soup.get_text(" ", strip=True))
    year_match = re.search(r"(20\d{2})年.*?" + re.escape(label.split("集計表")[0]), page_text)
    period_match = re.search(r"(\d)\s*[-〜]\s*(\d)月期", label)
    year = int(year_match.group(1)) if year_match else date.today().year
    period = f"{year}年{period_match.group(1)}-{period_match.group(2)}月期" if period_match else str(year)
    workbook = load_workbook(BytesIO(_get_excel(session, excel_url)), read_only=True, data_only=True)
    try:
        sheet = next((item for item in workbook.worksheets if _normal(item.title) == "参考1"), None)
        if sheet is None:
            raise RuntimeError("1人当たり旅行支出のシートを特定できませんでした。")
        value = sheet["F7"].value
        if not isinstance(value, (int, float)) or value <= 0:
            raise RuntimeError("1人当たり旅行支出の全国値が不正です。")
    finally:
        workbook.close()
    release_type = "第2次速報" if "2次速報" in label else "第1次速報"
    source = {
        "source_name": "観光庁「インバウンド消費動向調査」",
        "source_url": CONSUMPTION_PAGE,
        "excel_url": excel_url,
        "reference_period": period,
        "published_date": _published_from_page(response.text),
        "release_type": release_type,
    }
    indicator = {"value": round(float(value)), "unit": "円/人", "yoy_pct": None, "source_key": "inbound_consumption"}
    return source, indicator


def _lodging_statistics(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    inbound = _get(session, LAB_INBOUND_URL).json()
    analysis = _get(session, LAB_ANALYSIS_URL).json()
    metadata = inbound["metadata"]
    national = analysis["national"]
    source = {
        "source_name": "観光庁「宿泊旅行統計調査」",
        "source_url": metadata["source_url"],
        "excel_url": metadata["excel_url"],
        "reference_period": f'{metadata["year"]:04d}-{metadata["month"]:02d}',
        "published_date": metadata.get("updated_at"),
        "release_type": metadata.get("release_type"),
        "survey_scope": metadata.get("survey_scope"),
    }
    indicators = {
        "foreign_guest_nights": {
            "value": inbound["national"]["foreign_guest_nights"], "unit": "人泊", "yoy_pct": None,
            "source_key": "lodging_statistics",
        },
        "total_guest_nights_12m": {
            "value": national["total_guest_nights"], "unit": "人泊",
            "yoy_pct": round(float(national["demand_growth_rate"]) * 100, 4),
            "source_key": "lodging_statistics",
        },
        "occupancy_rate": {
            "value": national["occupancy_rate"], "unit": "%", "yoy_pct": None,
            "source_key": "lodging_statistics",
        },
    }
    return source, indicators


def validate_public_statistics(payload: dict[str, Any]) -> None:
    if payload.get("schema_version") != 1:
        raise ValueError("公的統計JSONのschema_versionが不正です。")
    sources = payload.get("sources")
    indicators = payload.get("indicators")
    if not isinstance(sources, dict) or not sources:
        raise ValueError("公的統計の出典がありません。")
    if not isinstance(indicators, dict) or not indicators:
        raise ValueError("公的統計の指標がありません。")
    for key, item in indicators.items():
        if item.get("value") is None or not isinstance(item.get("value"), (int, float)):
            raise ValueError(f"公的統計値が不正です: {key}")
        if item.get("source_key") not in sources:
            raise ValueError(f"公的統計の出典参照が不正です: {key}")


def refresh_public_statistics(cache_path: Path, session: requests.Session | None = None) -> dict[str, Any]:
    session = session or requests.Session()
    session.headers.update({"User-Agent": "TourismMarketSignal/1.0 (+https://lab.ugatta-llc.com/)"})
    cached: dict[str, Any] = {}
    if cache_path.exists():
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            validate_public_statistics(cached)
        except (OSError, json.JSONDecodeError, ValueError):
            cached = {}
    sources: dict[str, Any] = {}
    indicators: dict[str, Any] = {}
    errors: list[str] = []
    retrieved_at = datetime.now(JST).isoformat(timespec="seconds")
    for key, loader in (
        ("jnto_visitors", _jnto_statistics),
        ("lodging_statistics", _lodging_statistics),
        ("inbound_consumption", _consumption_statistics),
    ):
        try:
            source, values = loader(session)
            source.setdefault("retrieved_at", retrieved_at)
            sources[key] = source
            if key == "lodging_statistics":
                indicators.update(values)
            else:
                indicators["visitor_arrivals" if key == "jnto_visitors" else "spend_per_visitor"] = values
        except Exception as exc:
            logger.warning("公的統計を更新できませんでした: %s (%s)", key, exc)
            errors.append(key)
            cached_source = (cached.get("sources") or {}).get(key)
            if cached_source:
                sources[key] = cached_source
                indicator_keys = {
                    "jnto_visitors": ("visitor_arrivals",),
                    "lodging_statistics": ("foreign_guest_nights", "total_guest_nights_12m", "occupancy_rate"),
                    "inbound_consumption": ("spend_per_visitor",),
                }[key]
                for indicator_key in indicator_keys:
                    cached_indicator = (cached.get("indicators") or {}).get(indicator_key)
                    if cached_indicator:
                        indicators[indicator_key] = cached_indicator
    payload = {
        "schema_version": 1,
        "generated_at": retrieved_at,
        "sources": sources,
        "indicators": indicators,
        "comparability": {
            "lodging_2026_boundary": True,
            "lodging_growth_weight_multiplier": 0.5,
            "note": "宿泊旅行統計調査は2026年1月から集計・層化基準が変更されています。2025年以前との比較には、この変更の影響が含まれる可能性があります。",
        },
        "update_errors": errors,
        "is_stale": bool(errors),
    }
    if sources:
        validate_public_statistics(payload)
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return payload
    if cached:
        cached["is_stale"] = True
        cached["update_errors"] = errors
        return cached
    return {
        "schema_version": 1, "generated_at": None, "sources": {}, "indicators": {},
        "comparability": {}, "update_errors": errors, "is_stale": True,
    }
